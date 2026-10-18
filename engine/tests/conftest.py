from __future__ import annotations

import hashlib
import io
from pathlib import Path

import numpy as np
import pytest
from PIL import Image, ImageDraw, ImageFont


def render_text_image(lines: list[str], size=(1400, 500)) -> Image.Image:
    image = Image.new("RGB", size, "white")
    draw = ImageDraw.Draw(image)
    try:
        font = ImageFont.truetype("arial.ttf", 48)
    except OSError:
        font = ImageFont.load_default(size=48)
    y = 40
    for line in lines:
        draw.text((40, y), line, fill="black", font=font)
        y += 90
    return image


@pytest.fixture
def docs(tmp_path: Path) -> Path:
    """A folder with one document per supported family."""
    import docx
    import openpyxl
    from pptx import Presentation
    from reportlab.lib.pagesizes import A4
    from reportlab.pdfgen import canvas

    root = tmp_path / "raw"
    root.mkdir()

    d = docx.Document()
    d.core_properties.author = "Mario Rossi"
    d.core_properties.title = "Contratto di fornitura"
    d.add_heading("Contratto di fornitura", 1)
    d.add_paragraph("Il fornitore consegna 500 pallet entro il 31 marzo 2025 presso il magazzino di Bologna.")
    d.add_heading("Penali", 2)
    d.add_paragraph("In caso di ritardo si applica una penale del 2% per settimana.")
    table = d.add_table(rows=2, cols=2)
    table.cell(0, 0).text, table.cell(0, 1).text = "Voce", "Importo"
    table.cell(1, 0).text, table.cell(1, 1).text = "Pallet", "12.000 EUR"
    d.save(root / "contratto.docx")

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Budget"
    ws.append(["Reparto", "Spesa"])
    ws.append(["Marketing", 42000])
    ws.append(["Ricerca", 87000])
    wb.properties.creator = "Ufficio Finanza"
    wb.save(root / "budget.xlsx")

    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[1])
    slide.shapes.title.text = "Roadmap 2025"
    slide.placeholders[1].text = "Lancio del prodotto Aurora nel secondo trimestre"
    deck.save(root / "roadmap.pptx")

    c = canvas.Canvas(str(root / "manuale.pdf"), pagesize=A4)
    c.setTitle("Manuale operativo")
    c.setAuthor("Team Qualita")
    c.setFont("Helvetica-Bold", 20)
    c.drawString(72, 780, "Procedura di sicurezza")
    c.setFont("Helvetica", 11)
    c.drawString(72, 750, "Indossare sempre i guanti protettivi durante la manutenzione della pressa.")
    c.showPage()
    c.save()

    scan = render_text_image(["INVOICE NUMBER 4471", "TOTAL AMOUNT 980 EUR"])
    scan.save(root / "scansione.pdf", "PDF", resolution=150)
    render_text_image(["WAREHOUSE INVENTORY", "STEEL BOLTS 1200"]).save(root / "foto.png")

    (root / "note.md").write_text("# Note riunione\n\nDecisione: migrare il CRM entro giugno.\n\n## Azioni\n\nPreparare il piano.", encoding="utf-8")
    (root / "dati.csv").write_text("citta;abitanti\nRoma;2750000\nMilano;1350000\n", encoding="utf-8")
    (root / "config.json").write_text('{"server": "alpha", "porta": 8080}', encoding="utf-8")
    (root / "leggimi.txt").write_text("Questo archivio contiene i documenti del progetto Aurora.", encoding="utf-8")
    (root / "~$contratto.docx").write_bytes(b"lock file")  # Office lock file: must be ignored
    return root


def fake_vector(text: str, dim: int = 64) -> list[float]:
    """Deterministic bag-of-words embedding: shared words -> similar vectors."""
    vec = np.zeros(dim, dtype=np.float32)
    for word in text.lower().split():
        vec[int(hashlib.md5(word.encode()).hexdigest(), 16) % dim] += 1.0
    norm = np.linalg.norm(vec)
    return (vec / norm if norm else vec + 1e-3).tolist()


@pytest.fixture(autouse=True)
def no_model_preload(monkeypatch, request):
    """Unit tests must not ask a real LM Studio to load (fake) models."""
    if request.node.get_closest_marker("live"):
        return
    from vault_engine import app, pipeline
    from vault_engine.agent import runner

    async def noop(endpoints, refs):
        return None

    for module in (app, pipeline, runner):
        monkeypatch.setattr(module, "ensure_loaded", noop)


@pytest.fixture
def fake_embeddings(monkeypatch):
    from vault_engine import providers

    calls = {"texts": 0}

    async def embed(self, ref, texts):
        calls["texts"] += len(texts)
        return [fake_vector(t) for t in texts]

    monkeypatch.setattr(providers.ProviderClient, "embed", embed)
    return calls


@pytest.fixture
def png_bytes():
    def make(lines):
        buf = io.BytesIO()
        render_text_image(lines).save(buf, "PNG")
        return buf.getvalue()

    return make

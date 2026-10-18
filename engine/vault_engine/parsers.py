"""Format-aware parsing into ordered text blocks plus document metadata."""

from __future__ import annotations

import csv
import io
import json
import re
import struct
from dataclasses import dataclass, field
from datetime import datetime, timezone
from pathlib import Path
from typing import Callable, Literal

from PIL import Image, ImageSequence

from .ocr import OcrEngine, to_png

BlockKind = Literal["heading", "paragraph", "table", "ocr", "note"]

OFFICE_EXTS = {".docx", ".xlsx", ".pptx", ".msg"}
PDF_EXTS = {".pdf"}
IMAGE_EXTS = {".png", ".jpg", ".jpeg", ".tif", ".tiff", ".bmp", ".webp"}
TEXT_EXTS = {".txt", ".md", ".json", ".csv"}
SUPPORTED_EXTS = OFFICE_EXTS | PDF_EXTS | IMAGE_EXTS | TEXT_EXTS

# Rows per table block for spreadsheets/CSV, so one sheet never becomes one giant block.
TABLE_ROWS_PER_BLOCK = 40


@dataclass
class Block:
    text: str
    kind: BlockKind = "paragraph"
    page: int | None = None
    level: int = 0  # heading depth (1 = top)


@dataclass
class ParsedDocument:
    blocks: list[Block] = field(default_factory=list)
    title: str | None = None
    author: str | None = None
    created: str | None = None
    modified: str | None = None
    pages: int | None = None
    ocr_pages: int = 0
    ocr_engine: str | None = None
    warnings: list[str] = field(default_factory=list)

    @property
    def text(self) -> str:
        return "\n\n".join(b.text for b in self.blocks if b.text.strip())


class UnsupportedFormat(ValueError):
    pass


def parse_file(path: Path, ocr: OcrEngine, *, min_text_chars_per_page: int = 40, parser: str = "builtin") -> ParsedDocument:
    ext = path.suffix.lower()
    if ext not in SUPPORTED_EXTS:
        raise UnsupportedFormat(ext)
    if parser == "docling" and ext in PDF_EXTS | {".docx", ".pptx", ".xlsx"}:
        return _parse_docling(path)
    return parse_bytes(path.read_bytes(), ext, ocr, min_text_chars_per_page=min_text_chars_per_page, name=path.name)


def parse_bytes(data: bytes, ext: str, ocr: OcrEngine, *, min_text_chars_per_page: int = 40, name: str = "") -> ParsedDocument:
    handler: dict[str, Callable[..., ParsedDocument]] = {
        ".pdf": lambda: _parse_pdf(data, ocr, min_text_chars_per_page),
        ".docx": lambda: _parse_docx(data),
        ".xlsx": lambda: _parse_xlsx(data),
        ".pptx": lambda: _parse_pptx(data),
        ".msg": lambda: _parse_msg(data, ocr, min_text_chars_per_page),
        ".md": lambda: _parse_markdown(_decode(data)),
        ".txt": lambda: _parse_plain(_decode(data)),
        ".json": lambda: _parse_json(_decode(data)),
        ".csv": lambda: _parse_csv(_decode(data)),
    }
    if ext in IMAGE_EXTS:
        return _parse_image(data, ocr)
    if ext not in handler:
        raise UnsupportedFormat(ext or name)
    return handler[ext]()


# --------------------------------------------------------------------------- text


def _decode(data: bytes) -> str:
    for encoding in ("utf-8-sig", "utf-16") if data[:2] in (b"\xff\xfe", b"\xfe\xff") else ("utf-8-sig",):
        try:
            return data.decode(encoding)
        except UnicodeDecodeError:
            pass
    return data.decode("cp1252", errors="replace")


def _paragraphs(text: str, page: int | None = None) -> list[Block]:
    return [Block(p.strip(), "paragraph", page) for p in re.split(r"\n\s*\n", text) if p.strip()]


def _parse_plain(text: str) -> ParsedDocument:
    return ParsedDocument(blocks=_paragraphs(text))


MD_HEADING = re.compile(r"^(#{1,6})\s+(.+?)\s*#*\s*$")


def _parse_markdown(text: str) -> ParsedDocument:
    blocks: list[Block] = []
    buf: list[str] = []
    in_fence = False

    def flush() -> None:
        if buf:
            blocks.extend(_paragraphs("\n".join(buf)))
            buf.clear()

    for line in text.splitlines():
        if line.lstrip().startswith("```"):
            in_fence = not in_fence
        match = None if in_fence else MD_HEADING.match(line)
        if match:
            flush()
            blocks.append(Block(match.group(2), "heading", level=len(match.group(1))))
        else:
            buf.append(line)
    flush()
    title = next((b.text for b in blocks if b.kind == "heading"), None)
    return ParsedDocument(blocks=blocks, title=title)


def _parse_json(text: str) -> ParsedDocument:
    try:
        pretty = json.dumps(json.loads(text), indent=2, ensure_ascii=False)
    except json.JSONDecodeError:
        pretty = text
    return ParsedDocument(blocks=_paragraphs(pretty) if "\n\n" in pretty else [Block(pretty)])


def _table_blocks(rows: list[list[str]], caption: str | None = None, page: int | None = None) -> list[Block]:
    """Markdown tables, split every TABLE_ROWS_PER_BLOCK rows with the header repeated."""
    rows = [[c.replace("\n", " ").strip() for c in r] for r in rows if any(c.strip() for c in r)]
    if not rows:
        return []
    width = max(len(r) for r in rows)
    rows = [r + [""] * (width - len(r)) for r in rows]
    header, body = rows[0], rows[1:] or [[""] * width]
    head = "| " + " | ".join(header) + " |\n|" + "---|" * width
    blocks = []
    for start in range(0, len(body), TABLE_ROWS_PER_BLOCK):
        part = body[start : start + TABLE_ROWS_PER_BLOCK]
        table = head + "\n" + "\n".join("| " + " | ".join(r) + " |" for r in part)
        if caption:
            table = f"{caption}\n{table}"
        blocks.append(Block(table, "table", page))
    return blocks


def _parse_csv(text: str) -> ParsedDocument:
    try:
        dialect = csv.Sniffer().sniff(text[:4096], delimiters=",;\t|")
    except csv.Error:
        dialect = csv.excel
    rows = list(csv.reader(io.StringIO(text), dialect))
    return ParsedDocument(blocks=_table_blocks(rows))


# --------------------------------------------------------------------------- images


def _parse_image(data: bytes, ocr: OcrEngine) -> ParsedDocument:
    doc = ParsedDocument(ocr_engine=ocr.name)
    with Image.open(io.BytesIO(data)) as image:
        frames = [frame.copy() for frame in ImageSequence.Iterator(image)]
    doc.pages = len(frames)
    for number, frame in enumerate(frames, start=1):
        result = ocr.recognize(to_png(frame))
        doc.ocr_pages += 1
        if result.text:
            doc.blocks.append(Block(result.text, "ocr", number if len(frames) > 1 else None))
    if not doc.blocks:
        doc.warnings.append("OCR found no text in image")
    return doc


# --------------------------------------------------------------------------- PDF


def _pdf_date(value: str | None) -> str | None:
    # PDF dates look like "D:20240131120000+01'00'".
    if not value:
        return None
    match = re.match(r"D?:?(\d{4})(\d{2})?(\d{2})?(\d{2})?(\d{2})?(\d{2})?", value)
    if not match:
        return None
    parts = [int(p) if p else d for p, d in zip(match.groups(), (0, 1, 1, 0, 0, 0))]
    try:
        return datetime(*parts, tzinfo=timezone.utc).isoformat()
    except ValueError:
        return None


def _parse_pdf(data: bytes, ocr: OcrEngine, min_chars: int) -> ParsedDocument:
    import pymupdf  # noqa: PLC0415

    doc = ParsedDocument(ocr_engine=None)
    with pymupdf.open(stream=data, filetype="pdf") as pdf:
        meta = pdf.metadata or {}
        doc.title = meta.get("title") or None
        doc.author = meta.get("author") or None
        doc.created = _pdf_date(meta.get("creationDate"))
        doc.modified = _pdf_date(meta.get("modDate"))
        doc.pages = pdf.page_count
        sizes = _pdf_font_sizes(pdf)
        body_size = sizes[len(sizes) // 2] if sizes else 11.0
        for page in pdf:
            number = page.number + 1
            text_chars = len(page.get_text("text").strip())
            if text_chars < min_chars:
                # Raster/scanned page: render at 200 dpi and OCR it.
                pix = page.get_pixmap(dpi=200)
                result = ocr.recognize(pix.tobytes("png"))
                doc.ocr_pages += 1
                doc.ocr_engine = result.engine
                if result.text:
                    doc.blocks.append(Block(result.text, "ocr", number))
                continue
            doc.blocks.extend(_pdf_page_blocks(page, number, body_size))
    return doc


def _pdf_font_sizes(pdf) -> list[float]:
    sizes: list[float] = []
    for page in list(pdf)[:20]:
        for block in page.get_text("dict")["blocks"]:
            for line in block.get("lines", []):
                for span in line["spans"]:
                    if span["text"].strip():
                        sizes.extend([round(span["size"], 1)] * len(span["text"]))
    return sorted(sizes)


def _pdf_page_blocks(page, number: int, body_size: float) -> list[Block]:
    # (top y, block) pairs so tables and text interleave in reading order.
    positioned: list[tuple[float, Block]] = []
    table_rects = []
    try:
        for table in page.find_tables().tables:
            rows = [[cell or "" for cell in row] for row in table.extract()]
            table_blocks = _table_blocks(rows, page=number)
            if table_blocks:
                table_rects.append(table.bbox)
                positioned.extend((table.bbox[1], b) for b in table_blocks)
    except Exception:  # noqa: BLE001 - table detection is best effort
        positioned, table_rects = [], []

    def in_table(bbox) -> bool:
        return any(bbox[0] >= r[0] - 2 and bbox[1] >= r[1] - 2 and bbox[2] <= r[2] + 2 and bbox[3] <= r[3] + 2 for r in table_rects)

    for block in page.get_text("dict", sort=True)["blocks"]:
        if block.get("type") != 0 or in_table(block["bbox"]):
            continue
        lines, sizes, bold = [], [], True
        for line in block["lines"]:
            text = "".join(span["text"] for span in line["spans"]).strip()
            if text:
                lines.append(text)
                sizes.extend(span["size"] for span in line["spans"] if span["text"].strip())
                bold = bold and all("bold" in span["font"].lower() or span["flags"] & 16 for span in line["spans"] if span["text"].strip())
        if not lines:
            continue
        text = _join_pdf_lines(lines)
        size = max(sizes) if sizes else body_size
        is_heading = len(text) < 120 and len(lines) <= 2 and (size >= body_size * 1.25 or (bold and size >= body_size))
        level = 1 if size >= body_size * 1.6 else 2 if size >= body_size * 1.25 else 3
        positioned.append((block["bbox"][1], Block(text, "heading" if is_heading else "paragraph", number, level if is_heading else 0)))
    positioned.sort(key=lambda item: item[0])
    return [b for _, b in positioned]


def _join_pdf_lines(lines: list[str]) -> str:
    out = lines[0]
    for line in lines[1:]:
        # Re-join words hyphenated across line breaks.
        if out.endswith("-") and line[:1].islower():
            out = out[:-1] + line
        else:
            out += " " + line
    return out


# --------------------------------------------------------------------------- Office


def _iso(value) -> str | None:
    if isinstance(value, datetime):
        return (value if value.tzinfo else value.replace(tzinfo=timezone.utc)).isoformat()
    return None


def _parse_docx(data: bytes) -> ParsedDocument:
    import docx  # noqa: PLC0415
    from docx.table import Table  # noqa: PLC0415
    from docx.text.paragraph import Paragraph  # noqa: PLC0415

    document = docx.Document(io.BytesIO(data))
    props = document.core_properties
    doc = ParsedDocument(
        title=props.title or None,
        author=props.author or None,
        created=_iso(props.created),
        modified=_iso(props.modified),
    )
    for child in document.element.body.iterchildren():
        tag = child.tag.rsplit("}", 1)[-1]
        if tag == "p":
            para = Paragraph(child, document)
            text = para.text.strip()
            if not text:
                continue
            style = (para.style.name if para.style is not None else "") or ""
            match = re.match(r"(?:Heading|Titolo)\s*(\d)", style)
            if style == "Title" or match:
                doc.blocks.append(Block(text, "heading", level=int(match.group(1)) if match else 1))
            else:
                doc.blocks.append(Block(text))
        elif tag == "tbl":
            table = Table(child, document)
            rows = [[cell.text for cell in row.cells] for row in table.rows]
            doc.blocks.extend(_table_blocks(rows))
    return doc


def _parse_xlsx(data: bytes) -> ParsedDocument:
    import openpyxl  # noqa: PLC0415

    workbook = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    props = workbook.properties
    doc = ParsedDocument(
        title=props.title or None,
        author=props.creator or None,
        created=_iso(props.created),
        modified=_iso(props.modified),
        pages=len(workbook.sheetnames),
    )
    for index, sheet in enumerate(workbook.worksheets, start=1):
        rows = [["" if v is None else str(v) for v in row] for row in sheet.iter_rows(values_only=True)]
        rows = [r for r in rows if any(c.strip() for c in r)]
        if not rows:
            continue
        doc.blocks.append(Block(f"Foglio: {sheet.title}", "heading", index, level=2))
        doc.blocks.extend(_table_blocks(rows, page=index))
    workbook.close()
    return doc


def _parse_pptx(data: bytes) -> ParsedDocument:
    from pptx import Presentation  # noqa: PLC0415

    deck = Presentation(io.BytesIO(data))
    props = deck.core_properties
    doc = ParsedDocument(
        title=props.title or None,
        author=props.author or None,
        created=_iso(props.created),
        modified=_iso(props.modified),
        pages=len(deck.slides),
    )
    for number, slide in enumerate(deck.slides, start=1):
        title_shape = slide.shapes.title
        title = title_shape.text_frame.text.strip() if title_shape is not None and title_shape.has_text_frame else ""
        doc.blocks.append(Block(f"Slide {number}" + (f": {title}" if title else ""), "heading", number, level=2))
        for shape in slide.shapes:
            if shape is title_shape:
                continue
            if shape.has_text_frame and shape.text_frame.text.strip():
                doc.blocks.append(Block(shape.text_frame.text.strip(), "paragraph", number))
            if getattr(shape, "has_table", False) and shape.has_table:
                rows = [[cell.text for cell in row.cells] for row in shape.table.rows]
                doc.blocks.extend(_table_blocks(rows, page=number))
        if slide.has_notes_slide:
            notes = slide.notes_slide.notes_text_frame.text.strip()
            if notes:
                doc.blocks.append(Block(f"Note: {notes}", "note", number))
    return doc


# --------------------------------------------------------------------------- Outlook .msg


def _msg_string(ole, stream_prefix: str, prop: str) -> str | None:
    for kind, encoding in (("001F", "utf-16-le"), ("001E", "cp1252")):
        name = f"{stream_prefix}__substg1.0_{prop}{kind}"
        if ole.exists(name):
            return ole.openstream(name).read().decode(encoding, errors="replace").rstrip("\x00").strip()
    return None


def _msg_filetime(ole, prop_tag: int) -> str | None:
    """Reads a PT_SYSTIME property from the top-level fixed-length property stream."""
    if not ole.exists("__properties_version1.0"):
        return None
    raw = ole.openstream("__properties_version1.0").read()
    for offset in range(32, len(raw) - 15, 16):
        tag, _flags, value = struct.unpack_from("<IIQ", raw, offset)
        if tag == prop_tag and value:
            return datetime.fromtimestamp((value - 116444736000000000) / 1e7, tz=timezone.utc).isoformat()
    return None


def _parse_msg(data: bytes, ocr: OcrEngine, min_chars: int) -> ParsedDocument:
    import olefile  # noqa: PLC0415

    ole = olefile.OleFileIO(io.BytesIO(data))
    try:
        subject = _msg_string(ole, "", "0037")
        sender = _msg_string(ole, "", "0C1A") or _msg_string(ole, "", "0042")
        sender_email = _msg_string(ole, "", "5D01") or _msg_string(ole, "", "0C1F")
        to = _msg_string(ole, "", "0E04")
        cc = _msg_string(ole, "", "0E03")
        body = _msg_string(ole, "", "1000") or ""
        sent = _msg_filetime(ole, 0x00390040) or _msg_filetime(ole, 0x0E060040)
        header = [f"Oggetto: {subject or ''}", f"Da: {sender or ''}" + (f" <{sender_email}>" if sender_email else "")]
        if to:
            header.append(f"A: {to}")
        if cc:
            header.append(f"Cc: {cc}")
        if sent:
            header.append(f"Data: {sent}")
        doc = ParsedDocument(title=subject, author=sender, created=sent, modified=sent)
        doc.blocks.append(Block("\n".join(header), "note"))
        doc.blocks.extend(_paragraphs(body))

        attachments = sorted({entry[0] for entry in ole.listdir() if entry[0].startswith("__attach_version1.0_")})
        for attach in attachments:
            prefix = attach + "/"
            filename = _msg_string(ole, prefix, "3707") or _msg_string(ole, prefix, "3001") or ""
            data_stream = prefix + "__substg1.0_37010102"
            ext = Path(filename).suffix.lower()
            if not filename or ext not in SUPPORTED_EXTS - {".msg"} or not ole.exists(data_stream):
                continue
            try:
                inner = parse_bytes(ole.openstream(data_stream).read(), ext, ocr, min_text_chars_per_page=min_chars, name=filename)
            except Exception as exc:  # noqa: BLE001 - one bad attachment must not fail the email
                doc.warnings.append(f"attachment {filename}: {exc}")
                continue
            doc.blocks.append(Block(f"Allegato: {filename}", "heading", level=2))
            doc.blocks.extend(inner.blocks)
            doc.ocr_pages += inner.ocr_pages
            doc.ocr_engine = doc.ocr_engine or inner.ocr_engine
        return doc
    finally:
        ole.close()


# --------------------------------------------------------------------------- Docling (optional)


def _parse_docling(path: Path) -> ParsedDocument:
    try:
        from docling.document_converter import DocumentConverter  # noqa: PLC0415
    except ImportError as exc:
        raise RuntimeError("Docling parser selected but not installed: uv sync --extra docling") from exc
    result = DocumentConverter().convert(str(path))
    doc = _parse_markdown(result.document.export_to_markdown())
    doc.warnings.append("parsed with Docling")
    return doc
